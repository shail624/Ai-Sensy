"""Excel/JSON import & export format tests (FR-CON-04, FR-CON-15) — Module 2 Step 5D."""

from __future__ import annotations

import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from app.core.config import settings
from app.crm.csv_io import EXPORT_COLUMNS, map_and_validate, read_csv
from app.crm.formats import read_xlsx
from app.models.job_records import STATUS_COMPLETED, STATUS_READY
from app.services.export_service import ExportService
from app.services.import_service import ImportService
from app.storage.base import get_provider

PASSWORD = "Sup3r-Secret-Pass1"
MAPPING = {"Phone": "phone_e164", "Name": "full_name"}
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _local_storage(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_local_path", str(tmp_path))
    monkeypatch.setattr(settings, "storage_backend", "local")
    monkeypatch.setattr(settings, "storage_public_base_url", "")
    import app.storage.local  # noqa: F401 - registers the provider


def _xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def _headers(client, email: str) -> dict[str, str]:
    resp = await client.post("/api/v1/auth/login", json={"email": email, "password": PASSWORD})
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


async def _owner(client, make_user) -> dict[str, str]:
    await make_user(email="owner@vi.co", password=PASSWORD, is_superuser=True)
    return await _headers(client, "owner@vi.co")


async def _upload_xlsx(client, h, data: bytes) -> str:
    resp = await client.post(
        "/api/v1/media/upload",
        headers=h,
        files={"file": ("contacts.xlsx", data, XLSX_MIME)},
        data={"media_type": "document"},
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


async def _seed(client, h) -> None:
    for phone, name in (("+14155550001", "Alice"), ("+14155550002", "Bob")):
        resp = await client.post(
            "/api/v1/contacts", headers=h, json={"phone_e164": phone, "full_name": name}
        )
        assert resp.status_code == 201, resp.text


# --- Reader (pure) -----------------------------------------------------------
def test_read_xlsx_matches_the_csv_row_shape() -> None:
    """The two readers are interchangeable, so mapping/validation stay format-agnostic."""
    data = _xlsx([["Phone", "Name"], ["+14155550001", "Alice"], ["+14155550002", "Bob"]])
    assert read_xlsx(data) == read_csv(b"Phone,Name\n+14155550001,Alice\n+14155550002,Bob\n")


def test_read_xlsx_renders_numeric_phone_without_scientific_notation() -> None:
    """A phone typed as a number must not arrive as '1.4155550001e+10'."""
    rows = read_xlsx(_xlsx([["Phone", "Name"], [14155550001, "Alice"]]))
    assert rows[0]["Phone"] == "14155550001"


def test_read_xlsx_skips_blank_rows_and_handles_empty_sheet() -> None:
    rows = read_xlsx(_xlsx([["Phone", "Name"], ["+14155550001", "Alice"], [None, None]]))
    assert len(rows) == 1
    assert read_xlsx(_xlsx([])) == []


def test_xlsx_rows_flow_through_the_existing_validation() -> None:
    data = _xlsx([["Phone", "Name"], ["+14155550001", "Alice"], ["not-a-phone", "Broken"]])
    result = map_and_validate(read_xlsx(data), MAPPING)
    assert [r.fields["full_name"] for r in result.rows] == ["Alice"]
    assert len(result.errors) == 1 and "E.164" in result.errors[0].error


# --- Excel import (FR-CON-04) ------------------------------------------------
async def test_xlsx_import_creates_contacts_and_reports_bad_rows(
    client, make_user, session_factory
) -> None:
    h = await _owner(client, make_user)
    upload = await _upload_xlsx(
        client,
        h,
        _xlsx([["Phone", "Name"], ["+14155550001", "Alice"], ["nope", "Broken"]]),
    )
    import app.crm.tasks as tasks

    tasks.run_contact_import.apply_async = lambda args, task_id: None
    resp = await client.post(
        "/api/v1/contacts/import",
        headers=h,
        json={"upload_id": upload, "format": "xlsx", "mapping": MAPPING},
    )
    assert resp.status_code == 202, resp.text
    import_id = resp.json()["job"]["id"]

    async with session_factory() as session:
        job = await ImportService(session).run(import_id)

    assert job.status == STATUS_COMPLETED
    assert (job.total_rows, job.success_rows, job.error_rows) == (2, 1, 1)

    listing = (await client.get("/api/v1/contacts", headers=h)).json()["data"]
    assert [c["full_name"] for c in listing] == ["Alice"]

    progress = (await client.get(f"/api/v1/contacts/import/{import_id}", headers=h)).json()
    assert progress["error_report_url"] is not None


async def test_import_rejects_unsupported_format(client, make_user) -> None:
    h = await _owner(client, make_user)
    upload = await _upload_xlsx(client, h, _xlsx([["Phone"], ["+14155550001"]]))
    resp = await client.post(
        "/api/v1/contacts/import",
        headers=h,
        json={"upload_id": upload, "format": "json", "mapping": MAPPING},
    )
    assert resp.status_code == 422, resp.text


# --- Excel / JSON export (FR-CON-15) ----------------------------------------
async def _start_export(client, h, file_format: str) -> str:
    import app.crm.tasks as tasks

    tasks.run_contact_export.apply_async = lambda args, task_id: None
    resp = await client.post(
        "/api/v1/contacts/export",
        headers=h,
        json={"format": file_format, "match_type": "all", "rules": []},
    )
    assert resp.status_code == 202, resp.text
    return resp.json()["job"]["id"]


async def test_xlsx_export_writes_a_workbook(client, make_user, session_factory) -> None:
    h = await _owner(client, make_user)
    await _seed(client, h)
    export_id = await _start_export(client, h, "xlsx")

    async with session_factory() as session:
        job = await ExportService(session).run(export_id)

    assert job.status == STATUS_READY and job.row_count == 2
    assert job.storage_key.endswith(".xlsx")

    data = await get_provider("local").get(job.storage_key)
    sheet = load_workbook(io.BytesIO(data), read_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == list(EXPORT_COLUMNS)
    assert {rows[1][2], rows[2][2]} == {"Alice", "Bob"}


async def test_json_export_writes_row_objects(client, make_user, session_factory) -> None:
    h = await _owner(client, make_user)
    await _seed(client, h)
    export_id = await _start_export(client, h, "json")

    async with session_factory() as session:
        job = await ExportService(session).run(export_id)

    assert job.status == STATUS_READY and job.row_count == 2
    assert job.storage_key.endswith(".json")

    rows = json.loads((await get_provider("local").get(job.storage_key)).decode())
    assert len(rows) == 2
    assert set(rows[0]) == set(EXPORT_COLUMNS)
    assert {r["full_name"] for r in rows} == {"Alice", "Bob"}


async def test_every_export_format_returns_the_same_rows(
    client, make_user, session_factory
) -> None:
    """Format is a rendering choice; the audience and columns are shared (FR-CON-15)."""
    h = await _owner(client, make_user)
    await _seed(client, h)
    counts = {}
    for file_format in ("csv", "xlsx", "json"):
        export_id = await _start_export(client, h, file_format)
        async with session_factory() as session:
            job = await ExportService(session).run(export_id)
        assert job.status == STATUS_READY
        counts[file_format] = job.row_count
    assert counts == {"csv": 2, "xlsx": 2, "json": 2}


async def test_csv_export_behaviour_is_unchanged(client, make_user, session_factory) -> None:
    """The CSV path still renders the same header + one line per contact."""
    h = await _owner(client, make_user)
    await _seed(client, h)
    export_id = await _start_export(client, h, "csv")

    async with session_factory() as session:
        job = await ExportService(session).run(export_id)

    assert job.storage_key.endswith(".csv")
    lines = (await get_provider("local").get(job.storage_key)).decode().splitlines()
    assert lines[0].split(",") == list(EXPORT_COLUMNS)
    assert len(lines) == 3


async def test_export_rejects_unsupported_format(client, make_user) -> None:
    h = await _owner(client, make_user)
    resp = await client.post(
        "/api/v1/contacts/export", headers=h, json={"format": "pdf", "match_type": "all"}
    )
    assert resp.status_code == 422, resp.text
